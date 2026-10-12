import io
from datetime import datetime
from typing import List, Dict, Any, Optional
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

class ReportService:
    def export_excel(self, records: List[Dict[str, Any]], filters: Optional[Dict[str, str]] = None) -> io.BytesIO:
        """
        Generate a styled Excel workbook (.xlsx) from attendance records.
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Davomat Hisoboti"

        # Brand Header
        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = "SmartCheck — Seamless Face ID Attendance Report"
        title_cell.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
        title_cell.fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 40

        # Filter Subheader
        ws.merge_cells("A2:H2")
        meta_text = f"Yaratilgan vaqt: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        if filters:
            filter_parts = [f"{k}: {v}" for k, v in filters.items() if v]
            if filter_parts:
                meta_text += f" | Filtrlar: {', '.join(filter_parts)}"
        ws["A2"].value = meta_text
        ws["A2"].font = Font(name="Arial", size=10, italic=True, color="475569")
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 20

        # Table Column Headers
        headers = [
            "№",
            "Talaba ID",
            "F.I.SH.",
            "Guruh",
            "Kurs nomi",
            "Sana & Vaqt",
            "Holat",
            "Aniqlik (Confidence)"
        ]
        ws.append([])  # blank row 3
        ws.append(headers)  # row 4

        header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1')
        )

        for col_idx in range(1, 9):
            cell = ws.cell(row=4, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[4].height = 26

        # Populate rows
        for idx, r in enumerate(records, start=1):
            ts = r.get("timestamp")
            if isinstance(ts, datetime):
                formatted_ts = ts.strftime("%Y-%m-%d %H:%M:%S")
            else:
                formatted_ts = str(ts)

            status_text = r.get("status", "").capitalize()
            conf_val = f"{float(r.get('confidence_score', 0.0)) * 100:.1f}%"

            row_data = [
                idx,
                r.get("student_id", ""),
                r.get("student_name", ""),
                r.get("student_group", ""),
                r.get("course_name", ""),
                formatted_ts,
                status_text,
                conf_val
            ]
            ws.append(row_data)
            current_row = ws.max_row
            ws.row_dimensions[current_row].height = 22

            # Style status cells with color codes
            for col_idx in range(1, 9):
                c = ws.cell(row=current_row, column=col_idx)
                c.font = Font(name="Arial", size=10)
                c.border = thin_border
                c.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 4, 7, 8] else "left", vertical="center")
                
                if col_idx == 7: # Status
                    if status_text.lower() == "present":
                        c.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                        c.font = Font(name="Arial", size=10, bold=True, color="166534")
                    elif status_text.lower() == "late":
                        c.fill = PatternFill(start_color="FEF9C3", end_color="FEF9C3", fill_type="solid")
                        c.font = Font(name="Arial", size=10, bold=True, color="854D0E")
                    else:
                        c.fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
                        c.font = Font(name="Arial", size=10, bold=True, color="991B1B")

        # Auto-fit column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def export_pdf(self, records: List[Dict[str, Any]], filters: Optional[Dict[str, str]] = None) -> io.BytesIO:
        """
        Generate a professional PDF document using ReportLab.
        """
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=landscape(letter),
            rightMargin=30,
            leftMargin=30,
            topMargin=30,
            bottomMargin=30
        )
        elements = []
        styles = getSampleStyleSheet()

        # Title
        title_style = ParagraphStyle(
            name="TitleStyle",
            parent=styles["Heading1"],
            fontSize=20,
            textColor=colors.HexColor("#1E3A8A"),
            spaceAfter=4,
            alignment=1
        )
        elements.append(Paragraph("SmartCheck — Attendance Report", title_style))

        # Subtitle
        sub_style = ParagraphStyle(
            name="SubStyle",
            parent=styles["Normal"],
            fontSize=10,
            textColor=colors.HexColor("#64748B"),
            spaceAfter=15,
            alignment=1
        )
        meta_str = f"Exported: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        if filters:
            parts = [f"{k}: {v}" for k, v in filters.items() if v]
            if parts:
                meta_str += f" | {', '.join(parts)}"
        elements.append(Paragraph(meta_str, sub_style))

        # Summary KPIs
        total = len(records)
        present = sum(1 for r in records if str(r.get("status", "")).lower() == "present")
        late = sum(1 for r in records if str(r.get("status", "")).lower() == "late")
        rate = round((present + late) / total * 100, 1) if total > 0 else 0.0

        kpi_data = [
            ["Total Records", "Present Count", "Late Count", "Attendance Rate"],
            [str(total), str(present), str(late), f"{rate}%"]
        ]
        kpi_table = Table(kpi_data, colWidths=[180, 180, 180, 180])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#F1F5F9")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor("#475569")),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
            ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor("#FFFFFF")),
            ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 1), (-1, 1), 14),
            ('TEXTCOLOR', (0, 1), (-1, 1), colors.HexColor("#1E293B")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        # Main Table
        table_headers = ["№", "Student ID", "Full Name", "Group", "Course", "Timestamp", "Status", "Confidence"]
        table_data = [table_headers]

        for idx, r in enumerate(records[:150], start=1):  # Cap at 150 records per PDF report
            ts = r.get("timestamp")
            if isinstance(ts, datetime):
                formatted_ts = ts.strftime("%Y-%m-%d %H:%M")
            else:
                formatted_ts = str(ts)[:16]

            conf_str = f"{float(r.get('confidence_score', 0.0)) * 100:.0f}%"
            table_data.append([
                str(idx),
                str(r.get("student_id", "")),
                str(r.get("student_name", ""))[:22],
                str(r.get("student_group", "")),
                str(r.get("course_name", ""))[:20],
                formatted_ts,
                str(r.get("status", "")).capitalize(),
                conf_str
            ])

        col_widths = [30, 75, 160, 75, 150, 110, 65, 60]
        attendance_table = Table(table_data, colWidths=col_widths, repeatRows=1)
        attendance_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2563EB")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('ALIGN', (0, 0), (0, -1), 'CENTER'),
            ('ALIGN', (5, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 9),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 7),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#FFFFFF"), colors.HexColor("#F8FAFC")]),
        ]))
        elements.append(attendance_table)

        doc.build(elements)
        buffer.seek(0)
        return buffer

report_service = ReportService()
